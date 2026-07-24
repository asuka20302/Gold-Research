"""Append the international-gold source migration to both project manuals."""

from __future__ import annotations

from pathlib import Path

from docx import Document


PROJECT = Path(__file__).resolve().parents[1]
ENGLISH_DOC = PROJECT / "docs" / "Gold_Model_Project_Notes.docx"
CHINESE_DOC = PROJECT / "docs" / "黄金量化预测项目完整流程.docx"


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def update_english() -> None:
    document = Document(ENGLISH_DOC)
    title = "International XAU/USD source migration - 2026-07-24"
    if any(paragraph.text.strip() == title for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_heading(title, level=1)
    document.add_paragraph(
        "The maintained prediction target has been switched from Shanghai "
        "Gold Exchange Au99.99 in CNY per gram to FXCM XAU/USD in USD per "
        "troy ounce. The old Chinese-gold code remains as a legacy comparison "
        "and is not used by the new international target loader."
    )
    document.add_heading("Source and target rules", level=2)
    add_bullet(
        document,
        "Preferred source: Tushare fx_daily, XAUUSD.FXCM, with GMT-dated "
        "bid and ask OHLC fields.",
    )
    add_bullet(
        document,
        "Preferred horizon-one label: log(next-session mid close / "
        "next-session mid open).",
    )
    add_bullet(
        document,
        "Current fallback: recover the previously cached Tushare/FXCM "
        "mid-close sequence and undo its old one-session Shanghai alignment.",
    )
    add_bullet(
        document,
        "Fallback label: log(next-session mid close / signal-session mid "
        "close). It is research-only because no next open or spread exists.",
    )
    add_bullet(
        document,
        "SGE price, Shanghai premium, and SHFE gold variables are excluded "
        "from the international target model.",
    )
    document.add_heading("Leakage controls and model comparison", level=2)
    document.add_paragraph(
        "The usable international sample is divided into 19 equal "
        "chronological blocks of 163 observations after excluding 17 oldest "
        "remainder rows. Blocks 1-9 select features using the existing "
        "p-value-constrained AIC rule. Block 10 selects model parameters by "
        "growth MAE. Blocks 11-18 estimate normalized inverse-growth-MAE "
        "voting weights. Blocks 10-18 fit the frozen final models, with "
        "unavailable end labels purged, and Block 19 is opened only for the "
        "common comparison."
    )
    document.add_paragraph(
        "The five candidates remain OLS, Ridge, Lasso, Elastic Net, and "
        "moving-block-bagged Huber. Directional accuracy is not an objective; "
        "the reported quantities are growth MAE/RMSE in percentage points and "
        "price MAE/RMSE in USD per troy ounce."
    )
    document.add_heading("First international result", level=2)
    document.add_paragraph(
        "The current Tushare token failed the FXCM endpoint smoke test, so "
        "the completed run used the verified close-only fallback. "
        "Training-only AIC retained seven factors: one-session XAU/USD "
        "return, five-session XAU/USD momentum, 20-session XAU/USD "
        "volatility, one-session oil return, one-session real-yield change, "
        "10-year breakeven inflation, and the GVZ level."
    )
    document.add_paragraph(
        "On Block 19, the zero-return baseline had the lowest growth MAE at "
        "1.2775 percentage points. Lasso and Elastic Net each had 1.2781; "
        "the weighted ensemble had 1.2807. The best model therefore did not "
        "beat the zero forecast. The price graph looks much closer than the "
        "return graph because each predicted next close starts from the "
        "known current close; that visual persistence is not forecast skill."
    )
    document.add_paragraph(
        "Decision: the data-source migration is complete, but the present "
        "international close-to-close model is not eligible for a trading "
        "P&L claim. The next data step is to restore a valid fx_daily token, "
        "rerun with --target-mode open_to_close, and evaluate bid/ask-aware "
        "execution before designing a strategy."
    )
    document.add_heading("Reproduction", level=2)
    document.add_paragraph(
        "Run: python international_gold_model.py --horizon 1"
    )
    document.add_paragraph(
        "Require full OHLC: python international_gold_model.py --horizon 1 "
        "--refresh --target-mode open_to_close"
    )
    document.save(ENGLISH_DOC)


def update_chinese() -> None:
    document = Document(CHINESE_DOC)
    title = "28. 国际黄金 XAU/USD 数据源切换（2026-07-24）"
    if any(paragraph.text.strip() == title for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_heading(title, level=1)
    document.add_paragraph(
        "维护中的预测目标已从上海黄金交易所 Au99.99（人民币/克）切换为 "
        "FXCM XAU/USD（美元/金衡盎司）。原中国黄金代码仅保留为历史对照，"
        "新国际黄金入口不再把它作为标签。"
    )
    document.add_heading("数据与标签规则", level=2)
    add_bullet(
        document,
        "首选数据：Tushare fx_daily 的 XAUUSD.FXCM，交易日期为 GMT，"
        "包含买卖盘开高低收。",
    )
    add_bullet(
        document,
        "完整 OHLC 下的一期标签：log（下一交易日中间收盘价 / "
        "下一交易日中间开盘价）。",
    )
    add_bullet(
        document,
        "当前 token 无法刷新接口，因此本次使用已缓存的 Tushare/FXCM "
        "中间收盘价，并撤销旧流程为上海收盘对齐所做的一期滞后。",
    )
    add_bullet(
        document,
        "收盘价回退标签：log（下一交易日中间收盘价 / 信号日中间收盘价）。"
        "因为没有下一日开盘和买卖价差，不能据此声称可执行 P&L。",
    )
    add_bullet(
        document,
        "国际模型排除 SGE 价格、上海溢价和 SHFE 黄金变量。",
    )
    document.add_heading("十九分组与模型", level=2)
    document.add_paragraph(
        "国际样本形成 19 个等长时间块，每块 163 条，最早 17 条余数不参与。"
        "第 1—9 块用于训练期 AIC 特征筛选，第 10 块按涨跌幅 MAE 选择超参数，"
        "第 11—18 块产生滚动样本外预测并计算 MAE 倒数权重，最终模型使用"
        "第 10—18 块训练，在预测第 19 块前删除尚不可观测的尾部标签。"
    )
    document.add_paragraph(
        "五个候选模型为 OLS、Ridge、Lasso、Elastic Net 和移动块 Bagged "
        "Huber。评价不使用方向准确率，而使用涨跌幅 MAE/RMSE（百分点）和"
        "价格 MAE/RMSE（美元/盎司）。"
    )
    document.add_heading("首轮结果与结论", level=2)
    document.add_paragraph(
        "训练期 AIC 保留 7 个变量：XAU/USD 一期收益、五期动量、20 期波动率、"
        "原油一期收益、实际利率一期变化、10 年盈亏平衡通胀率和 GVZ 水平。"
    )
    document.add_paragraph(
        "第 19 块中，零收益基准的涨跌幅 MAE 为 1.2775 个百分点；Lasso 与 "
        "Elastic Net 均为 1.2781；加权组合为 1.2807。因此当前规格没有"
        "战胜零预测。价格曲线看起来接近主要来自“下一日价格通常接近今日"
        "价格”的持久性，不能代替收益误差检验。"
    )
    document.add_paragraph(
        "结论：国际数据源切换已经完成，但当前收盘到收盘模型不具备交易 "
        "P&L 资格。下一步应恢复有效的 fx_daily 权限，以 "
        "--target-mode open_to_close 重新训练，并在买卖价差和实际成本下"
        "重新检验。"
    )
    document.save(CHINESE_DOC)


if __name__ == "__main__":
    update_english()
    update_chinese()
